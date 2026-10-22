import openpyxl, glob
p=glob.glob('/home/ubuntu/attachments/*/ITS_AND_RED_BUS_PNR*')[0]
wb=openpyxl.load_workbook(p, data_only=False)
wbv=openpyxl.load_workbook(p, data_only=True)
print(wb.sheetnames)
ws=wb['FINAL']; wsv=wbv['FINAL']
print('rows',ws.max_row,'cols',ws.max_column)
for row in ws.iter_rows(min_row=1, max_row=min(30,ws.max_row)):
    vals=[]
    for c in row:
        if c.value is not None:
            v=wsv[c.coordinate].value
            vals.append(f"{c.coordinate}={c.value!r}" + (f"→{v!r}" if str(c.value).startswith('=') else ""))
    if vals: print(vals)
